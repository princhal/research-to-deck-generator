"""Assembles a branded .pptx deck from a slide-plan JSON produced by the
Node/Claude synthesis step. Invoked as a subprocess from the BullMQ worker
(src/worker/processGenerate.ts) — this script does no network calls and
trusts its input has already been synthesized and validated upstream.

Usage: python3 generate_deck.py <input.json> <output.pptx>
"""

import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


def hex_to_rgb(hex_str: str) -> RGBColor:
    return RGBColor.from_string(hex_str)


def build_citation_map(citations: dict) -> tuple[dict, list]:
    """Assigns a stable citation number per unique paper, in first-appearance
    order of the chunk ids as they were defined in the input."""
    paper_order = []
    paper_by_id = {}
    chunk_to_paper_number = {}

    for chunk_id, meta in citations.items():
        paper_id = meta["paperId"]
        if paper_id not in paper_by_id:
            paper_by_id[paper_id] = meta
            paper_order.append(paper_id)
        number = paper_order.index(paper_id) + 1
        chunk_to_paper_number[str(chunk_id)] = number

    numbered_sources = [
        {"number": i + 1, **paper_by_id[pid]} for i, pid in enumerate(paper_order)
    ]
    return chunk_to_paper_number, numbered_sources


def add_title_slide(prs: Presentation, deck_title: str, subtitle: str, brand: dict):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = deck_title
    slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = hex_to_rgb(
        brand["colors"]["primary"]
    )
    slide.shapes.title.text_frame.paragraphs[0].font.name = brand["fonts"]["heading"]

    if len(slide.placeholders) > 1:
        subtitle_ph = slide.placeholders[1]
        subtitle_ph.text = subtitle
        subtitle_ph.text_frame.paragraphs[0].font.color.rgb = hex_to_rgb(
            brand["colors"]["muted"]
        )


def add_content_slide(prs: Presentation, slide_data: dict, chunk_to_paper_number: dict, brand: dict):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = slide_data["title"]
    slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = hex_to_rgb(
        brand["colors"]["primary"]
    )
    slide.shapes.title.text_frame.paragraphs[0].font.name = brand["fonts"]["heading"]

    body = slide.placeholders[1].text_frame
    body.clear()
    bullets = slide_data.get("bullets", [])
    for i, bullet in enumerate(bullets):
        paragraph = body.paragraphs[0] if i == 0 else body.add_paragraph()
        numbers = sorted(
            {
                chunk_to_paper_number.get(str(cid))
                for cid in bullet.get("sourceChunkIds", [])
                if chunk_to_paper_number.get(str(cid)) is not None
            }
        )
        citation_suffix = "".join(f"[{n}]" for n in numbers)
        paragraph.text = f"{bullet['text']} {citation_suffix}".strip()
        paragraph.font.size = Pt(18)
        paragraph.font.color.rgb = hex_to_rgb(brand["colors"]["text"])
        paragraph.font.name = brand["fonts"]["body"]

    notes = slide.notes_slide.notes_text_frame
    notes.text = slide_data.get("speakerNotes", "")


def add_sources_slide(prs: Presentation, numbered_sources: list, brand: dict):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Sources"
    slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = hex_to_rgb(
        brand["colors"]["primary"]
    )

    body = slide.placeholders[1].text_frame
    body.clear()
    for i, source in enumerate(numbered_sources):
        paragraph = body.paragraphs[0] if i == 0 else body.add_paragraph()
        authors = ", ".join(source.get("authors", [])) or "Unknown authors"
        year = source.get("year") or "n.d."
        line = f"[{source['number']}] {source['title']} — {authors} ({year})"
        if source.get("url"):
            line += f" — {source['url']}"
        paragraph.text = line
        paragraph.font.size = Pt(12)
        paragraph.font.color.rgb = hex_to_rgb(brand["colors"]["muted"])


def main():
    if len(sys.argv) != 3:
        print("Usage: python3 generate_deck.py <input.json> <output.pptx>", file=sys.stderr)
        sys.exit(1)

    input_path, output_path = Path(sys.argv[1]), Path(sys.argv[2])
    data = json.loads(input_path.read_text())

    brand = data["brand"]
    chunk_to_paper_number, numbered_sources = build_citation_map(data["citations"])

    prs = Presentation()
    add_title_slide(prs, data["deckTitle"], data.get("subtitle", ""), brand)
    for slide_data in data["slides"]:
        add_content_slide(prs, slide_data, chunk_to_paper_number, brand)
    add_sources_slide(prs, numbered_sources, brand)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    print(f"Deck written to {output_path}")


if __name__ == "__main__":
    main()
